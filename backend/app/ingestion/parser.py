import io
import os
import logging
from pathlib import Path
from typing import List, Dict, Any, Optional
from dataclasses import dataclass, field
from PIL import Image
from app.ingestion.preprocessor import clean_text
from app.ingestion.ocr import detect_ocr_needed, run_ocr_on_image

logger = logging.getLogger(__name__)


@dataclass
class PageContent:
    page_number: int
    text: str
    is_ocr_applied: bool = False
    word_count: int = 0


@dataclass
class ParsedDocument:
    filename: str
    mime_type: str
    pages: List[PageContent] = field(default_factory=list)
    full_text: str = ""
    ocr_applied: bool = False
    metadata: Dict[str, Any] = field(default_factory=dict)


def parse_pdf(file_path: str) -> ParsedDocument:
    pages: List[PageContent] = []
    ocr_applied_overall = False

    try:
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        total_pages = len(reader.pages)

        for idx, page in enumerate(reader.pages):
            page_num = idx + 1
            raw_text = page.extract_text() or ""
            cleaned = clean_text(raw_text)

            # Check if OCR is needed
            page_box = page.mediabox
            width = float(page_box.width) if page_box else 612.0
            height = float(page_box.height) if page_box else 792.0

            is_ocr = False
            if detect_ocr_needed(cleaned, width, height):
                # Attempt OCR on page images
                logger.info(f"Page {page_num} in {Path(file_path).name} requires OCR (text length={len(cleaned)})")
                ocr_text_accum = []
                try:
                    for img_obj in page.images:
                        ocr_txt = run_ocr_on_image(img_obj.data)
                        if ocr_txt:
                            ocr_text_accum.append(ocr_txt)
                except Exception as img_err:
                    logger.warning(f"Could not extract images from page {page_num}: {img_err}")

                if ocr_text_accum:
                    ocr_res = "\n".join(ocr_text_accum)
                    cleaned = clean_text(f"{cleaned}\n{ocr_res}")
                    is_ocr = True
                    ocr_applied_overall = True

            word_count = len(cleaned.split()) if cleaned else 0
            pages.append(PageContent(
                page_number=page_num,
                text=cleaned,
                is_ocr_applied=is_ocr,
                word_count=word_count
            ))

    except Exception as e:
        logger.error(f"Error parsing PDF with pypdf: {e}. Trying pdfplumber fallback.", exc_info=True)
        try:
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                for idx, page in enumerate(pdf.pages):
                    page_num = idx + 1
                    raw_text = page.extract_text() or ""
                    cleaned = clean_text(raw_text)
                    pages.append(PageContent(
                        page_number=page_num,
                        text=cleaned,
                        is_ocr_applied=False,
                        word_count=len(cleaned.split())
                    ))
        except Exception as e2:
            logger.error(f"pdfplumber also failed for {file_path}: {e2}")

    full_text = "\n\n".join(p.text for p in pages if p.text)
    return ParsedDocument(
        filename=Path(file_path).name,
        mime_type="application/pdf",
        pages=pages,
        full_text=full_text,
        ocr_applied=ocr_applied_overall,
        metadata={"total_pages": len(pages)}
    )


def parse_docx(file_path: str) -> ParsedDocument:
    pages: List[PageContent] = []
    text_blocks = []

    try:
        import docx
        doc = docx.Document(file_path)

        for p in doc.paragraphs:
            if p.text and p.text.strip():
                text_blocks.append(p.text.strip())

        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_blocks.append(row_text)

        full_text = clean_text("\n\n".join(text_blocks))
        pages.append(PageContent(
            page_number=1,
            text=full_text,
            is_ocr_applied=False,
            word_count=len(full_text.split())
        ))
    except Exception as e:
        logger.error(f"Error parsing DOCX file {file_path}: {e}", exc_info=True)

    return ParsedDocument(
        filename=Path(file_path).name,
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        pages=pages,
        full_text=pages[0].text if pages else "",
        ocr_applied=False,
        metadata={"paragraph_count": len(text_blocks)}
    )


def parse_text(file_path: str) -> ParsedDocument:
    text = ""
    for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
        try:
            with open(file_path, "r", encoding=encoding) as f:
                text = f.read()
                break
        except UnicodeDecodeError:
            continue

    cleaned = clean_text(text)
    page = PageContent(
        page_number=1,
        text=cleaned,
        is_ocr_applied=False,
        word_count=len(cleaned.split())
    )

    return ParsedDocument(
        filename=Path(file_path).name,
        mime_type="text/plain",
        pages=[page],
        full_text=cleaned,
        ocr_applied=False,
        metadata={"char_count": len(cleaned)}
    )


def parse_image(file_path: str) -> ParsedDocument:
    ocr_text = run_ocr_on_image(file_path)
    cleaned = clean_text(ocr_text)

    page = PageContent(
        page_number=1,
        text=cleaned,
        is_ocr_applied=True,
        word_count=len(cleaned.split())
    )

    ext = Path(file_path).suffix.lower()
    mime = f"image/{ext.replace('.', '')}"

    return ParsedDocument(
        filename=Path(file_path).name,
        mime_type=mime,
        pages=[page],
        full_text=cleaned,
        ocr_applied=True,
        metadata={"source": "ocr"}
    )


def parse_document(file_path: str) -> ParsedDocument:
    ext = Path(file_path).suffix.lower()
    logger.info(f"Parsing document: {file_path} (format={ext})")

    if ext == ".pdf":
        return parse_pdf(file_path)
    elif ext in [".docx", ".doc"]:
        return parse_docx(file_path)
    elif ext in [".txt", ".md", ".csv", ".log"]:
        return parse_text(file_path)
    elif ext in [".png", ".jpg", ".jpeg", ".tiff", ".bmp", ".webp"]:
        return parse_image(file_path)
    else:
        # Fallback to plain text read
        return parse_text(file_path)
